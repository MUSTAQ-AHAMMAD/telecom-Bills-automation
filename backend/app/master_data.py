"""
Imports the "TELECOMMUNICATION MASTER" spreadsheet — the real-world
inventory of every branch/line per entity+carrier, including which city
(Location) and branch each account number belongs to — into a local
SQLite database, and provides lookups used to attach Location/Branch to
scraped accounts and downloaded invoices.

The spreadsheet has one sheet per entity+carrier pair (e.g. "STC IBQ",
"Mobily MATCH", "Zain"), each with its own slightly different column set,
plus a "Closed" sheet (cancelled/closed lines, no entity/carrier per row)
and a "Sheet7" with no header row that doesn't fit this schema and is
skipped entirely.
"""
import os
import re
import sqlite3
from datetime import datetime
from typing import Optional

DB_PATH = os.path.join(os.path.dirname(__file__), "..", "data", "master.db")

# Maps each spreadsheet tab to the (entity, carrier) pair it represents.
# "Zain" (no suffix) sits with the IBQ group in the workbook and has no
# other entity marker, so it's inferred to be Ibraq's Zain account.
SHEET_TO_ENTITY_CARRIER = {
    "STC IBQ": ("Ibraq", "STC"),
    "Mobily IBQ": ("Ibraq", "Mobily"),
    "Zain": ("Ibraq", "Zain"),
    "STC MATCH ": ("Match", "STC"),
    "STC MATCH": ("Match", "STC"),
    "Mobily MATCH": ("Match", "Mobily"),
    "Zain Match": ("Match", "Zain"),
    "STC SALFA": ("Salfa", "STC"),
    "Mobily Salfa": ("Salfa", "Mobily"),
    "STC FEELIN": ("Feelin", "STC"),
    "Mobily Feelin": ("Feelin", "Mobily"),
}

# Sheets that don't map to a single entity+carrier, or don't fit the
# Location/Account Number schema at all.
CLOSED_SHEET = "Closed"
SKIPPED_SHEETS = {"Sheet7"}

# Column header variants (lowercased, stripped) -> canonical field name.
HEADER_ALIASES = {
    "location": "location",
    "cr": "cr",
    "conntract no": "cr",
    "contract no": "cr",
    "branch/users": "branch",
    "branch/user": "branch",
    "dept": "department",
    "department": "department",
    "department / code": "department",
    "connection": "connection",
    "account number": "account_number",
    "contract/account number": "account_number",
    "service number": "service_number",
    "service number (san)": "service_number",
    "service number (msisdn)": "service_number",
    "serial number": "serial_number",
    "s\\n": "serial_number",
    "stats": "status",
    "status": "status",
    "package": "package",
    "notes": "notes",
    "date": "closed_date",
}


def _normalize_header(value) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip().lower()
    text = re.sub(r"\s+", " ", text)
    return HEADER_ALIASES.get(text)


def _normalize_account_number(value) -> Optional[str]:
    """Excel stores account numbers as int/float/str inconsistently; normalize to a plain string."""
    if value is None:
        return None
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    text = str(value).strip()
    return text or None


def _find_header_row(ws, max_scan: int = 6) -> Optional[int]:
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1):
        cells = [str(c).strip().lower() if c is not None else "" for c in row]
        if "location" in cells:
            return row_idx
    return None


def _connect() -> sqlite3.Connection:
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    return sqlite3.connect(DB_PATH)


def _ensure_schema(conn: sqlite3.Connection) -> None:
    """
    Creates the tables if they don't exist yet, and migrates older databases
    forward (adding columns introduced later) — never drops existing data.
    Unlike the old `_init_schema`, this is safe to call before every
    operation, since re-importing the spreadsheet must not wipe manually
    added rows or rows a person has since hand-edited (see
    `import_master_excel` / `is_manually_edited`).
    """
    conn.execute(
        """CREATE TABLE IF NOT EXISTS master_accounts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            entity TEXT,
            carrier TEXT,
            sheet_name TEXT,
            location TEXT,
            cr TEXT,
            branch TEXT,
            department TEXT,
            connection TEXT,
            account_number TEXT,
            service_number TEXT,
            serial_number TEXT,
            status TEXT,
            package TEXT,
            notes TEXT
        )"""
    )
    conn.execute("CREATE INDEX IF NOT EXISTS idx_master_account_number ON master_accounts(carrier, account_number)")

    conn.execute(
        """CREATE TABLE IF NOT EXISTS closed_accounts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            location TEXT,
            cr TEXT,
            branch TEXT,
            connection TEXT,
            account_number TEXT,
            service_number TEXT,
            serial_number TEXT,
            status TEXT,
            closed_date TEXT
        )"""
    )
    conn.execute("CREATE INDEX IF NOT EXISTS idx_closed_account_number ON closed_accounts(account_number)")

    # Per-line bill charges for one month (from an "All stores Bills" sheet),
    # used to split an account's bill across cost centers — see
    # `bill_line_allocation`. STC's PDF bills can't be read as text, so the
    # sheet is the source of per-line amounts.
    conn.execute(
        """CREATE TABLE IF NOT EXISTS bill_lines (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            month TEXT, entity TEXT, carrier TEXT, account_number TEXT,
            service_number TEXT, store TEXT, location TEXT, cost_center TEXT, amount REAL,
            master_id INTEGER, sheet_row INTEGER
        )"""
    )
    conn.execute("CREATE INDEX IF NOT EXISTS idx_bill_lines_account ON bill_lines(entity, carrier, account_number, month)")

    existing_cols = {row[1] for row in conn.execute("PRAGMA table_info(master_accounts)")}
    if "cost_center" not in existing_cols:
        conn.execute("ALTER TABLE master_accounts ADD COLUMN cost_center TEXT")
    if "is_manually_edited" not in existing_cols:
        # Marks a row as protected from being wiped/overwritten by a future
        # spreadsheet re-import — set whenever a row is touched through the
        # CRUD API (create/update) or through a targeted enrichment pass
        # like `import_store_sheet`, so a person's correction (or an
        # enrichment from a different source file) survives the next
        # `import_master_excel` run.
        conn.execute("ALTER TABLE master_accounts ADD COLUMN is_manually_edited INTEGER DEFAULT 0")
    if "dept_code" not in existing_cols:
        conn.execute("ALTER TABLE master_accounts ADD COLUMN dept_code TEXT")


def import_master_excel(path: str) -> dict:
    """
    Refreshes the local master database from the spreadsheet at `path` —
    the source of truth for the branch/location inventory. Only rows this
    same automated import created are replaced (matched by sheet_name !=
    "Manual Entry" and not since hand-edited); rows added via the Master
    Data CRUD screen, or spreadsheet rows a person has since corrected
    (`is_manually_edited`), are left alone rather than wiped. Returns a
    per-sheet row-count summary.
    """
    import openpyxl  # imported lazily: only needed for this import step

    if not os.path.isfile(path):
        raise FileNotFoundError(f"Master spreadsheet not found: {path}")

    wb = openpyxl.load_workbook(path, data_only=True)
    if _flat_header_map(wb.worksheets[0]) is not None:
        # A single flat sheet with its own Entity/Carrier columns (the
        # Master Data table's own layout) — merged in, never wiped.
        return _import_flat_master(wb.worksheets[0])

    conn = _connect()
    summary = {}
    try:
        _ensure_schema(conn)
        conn.execute(
            "DELETE FROM master_accounts WHERE sheet_name != 'Manual Entry' "
            "AND (is_manually_edited IS NULL OR is_manually_edited = 0)"
        )
        conn.execute("DELETE FROM closed_accounts")

        # Whatever's left at this point is protected (manually edited, or a
        # manual entry) — re-inserting a fresh copy of the same logical row
        # from the sheet would duplicate it, so remember its natural key
        # (sheet_name, account_number, branch) and skip re-inserting a match.
        # NOTE: a handful of accounts have several service lines under the
        # identical branch name (same key) — editing one such row protects
        # that key and so skips re-inserting its siblings too. Not an issue
        # for anything edited so far (verified against real data), but a
        # true fix would need service_number in the key as well.
        protected_keys = {
            (row[0], row[1], row[2])
            for row in conn.execute("SELECT sheet_name, account_number, branch FROM master_accounts")
        }

        for sheet_name in wb.sheetnames:
            if sheet_name in SKIPPED_SHEETS:
                summary[sheet_name] = {"skipped": True}
                continue

            ws = wb[sheet_name]
            header_row = _find_header_row(ws)
            if header_row is None:
                summary[sheet_name] = {"skipped": True, "reason": "no header row found"}
                continue

            headers = [
                _normalize_header(c)
                for c in next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
            ]

            is_closed = sheet_name == CLOSED_SHEET
            entity_carrier = SHEET_TO_ENTITY_CARRIER.get(sheet_name)
            if not is_closed and not entity_carrier:
                summary[sheet_name] = {"skipped": True, "reason": "no entity/carrier mapping configured"}
                continue

            inserted = 0
            preserved = 0
            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                record = {}
                for header, value in zip(headers, row):
                    if header and value is not None:
                        record[header] = value

                account_number = _normalize_account_number(record.get("account_number"))
                # A row with no account number and no branch name is almost
                # certainly a blank trailing row — skip it.
                if not account_number and not record.get("branch"):
                    continue

                def clean(field):
                    v = record.get(field)
                    if v is None:
                        return None
                    text = " ".join(str(v).strip().split())
                    if field == "location" and text:
                        # Safe casing-only normalization (DAMMAM/Dammam/dammam
                        # -> Dammam) so identical cities aren't split into
                        # separate location-stat buckets. Genuinely different
                        # spellings (Jazan/Jizan, Madina/Madinah, Ahsa/Hassa)
                        # are left alone rather than guessing they're the same place.
                        text = text.title()
                    return text or None

                if is_closed:
                    conn.execute(
                        """INSERT INTO closed_accounts
                           (location, cr, branch, connection, account_number, service_number, serial_number, status, closed_date)
                           VALUES (?,?,?,?,?,?,?,?,?)""",
                        (
                            clean("location"), clean("cr"), clean("branch"), clean("connection"),
                            account_number, clean("service_number"), clean("serial_number"),
                            clean("status"), clean("closed_date"),
                        ),
                    )
                else:
                    entity, carrier = entity_carrier
                    if (sheet_name, account_number, clean("branch")) in protected_keys:
                        # Already present as a protected row (manually edited
                        # since the last import) — don't insert a duplicate.
                        preserved += 1
                        continue
                    conn.execute(
                        """INSERT INTO master_accounts
                           (entity, carrier, sheet_name, location, cr, branch, department, connection,
                            account_number, service_number, serial_number, status, package, notes)
                           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                        (
                            entity, carrier, sheet_name, clean("location"), clean("cr"), clean("branch"),
                            clean("department"), clean("connection"), account_number,
                            clean("service_number"), clean("serial_number"), clean("status"),
                            clean("package"), clean("notes"),
                        ),
                    )
                inserted += 1

            summary[sheet_name] = {
                "entity": entity_carrier[0] if entity_carrier else None,
                "carrier": entity_carrier[1] if entity_carrier else None,
                "rows_imported": inserted,
                "rows_preserved": preserved,
            }

        conn.commit()
    finally:
        conn.close()

    return summary


# Flat layout (one sheet, one row per service line, with its own Entity and
# Carrier columns — same columns as the Master Data table/export).
FLAT_HEADER_ALIASES = {
    **HEADER_ALIASES,
    "entity": "entity",
    "carrier": "carrier",
    "branch / store": "branch",
    "branch/store": "branch",
    "cost center": "cost_center",
    "dept code": "dept_code",
}
FLAT_SHEET_NAME = "Master Upload"
COST_CENTER_WIDTH = 4


def _flat_header_map(ws) -> Optional[dict]:
    """{canonical field: column index} if row 1 is a flat master header (has Entity, Carrier and Service Number), else None."""
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    mapping = {}
    for idx, value in enumerate(header):
        if value is None:
            continue
        field = FLAT_HEADER_ALIASES.get(re.sub(r"\s+", " ", str(value).strip().lower()))
        if field and field not in mapping:
            mapping[field] = idx
    if {"entity", "carrier", "service_number"} <= mapping.keys():
        return mapping
    return None


def _flat_cell_text(field: str, value) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = " ".join(str(value).strip().split())
    if not text:
        return None
    if field == "location":
        text = text.title()
    elif field == "cost_center" and text.isdigit():
        # Excel drops leading zeros (0802 -> 802); cost centers are fixed-width GL segments.
        text = text.zfill(COST_CENTER_WIDTH)
    return text


def _import_flat_master(ws) -> dict:
    """
    Merges a flat master sheet into master_accounts, matching each row to an
    existing line by (entity, carrier, service_number): matches are updated
    in place, unknown lines are added. Lines not in the sheet are left
    untouched. Every touched row is flagged `is_manually_edited` so a later
    per-carrier workbook import won't wipe it.
    """
    columns = _flat_header_map(ws)
    fields = [f for f in MASTER_ACCOUNT_FIELDS if f in columns]
    conn = _connect()
    summary: dict = {}
    try:
        _ensure_schema(conn)
        for row in ws.iter_rows(min_row=2, values_only=True):
            record = {
                f: (_normalize_account_number(row[columns[f]]) if f in ("account_number", "service_number")
                    else _flat_cell_text(f, row[columns[f]]))
                for f in fields if columns[f] < len(row)
            }
            entity, carrier, service = record.get("entity"), record.get("carrier"), record.get("service_number")
            if not (entity and carrier and service):
                continue
            stats = summary.setdefault(f"{ws.title} — {entity} {carrier}", {
                "entity": entity, "carrier": carrier, "rows_imported": 0,
                "rows_updated": 0, "rows_added": 0, "rows_preserved": 0,
            })

            existing = conn.execute(
                "SELECT id, branch FROM master_accounts WHERE entity = ? AND carrier = ? AND service_number = ?",
                (entity, carrier, service),
            ).fetchall()
            if existing:
                for row_id, old_branch in existing:
                    values = dict(record)
                    # Same leading-zero loss for numeric branch codes (0908 -> 908): keep the stored spelling.
                    if (old_branch and values.get("branch") and old_branch.isdigit()
                            and values["branch"].isdigit() and int(old_branch) == int(values["branch"])):
                        values["branch"] = old_branch
                    values["is_manually_edited"] = 1
                    conn.execute(
                        f"UPDATE master_accounts SET {', '.join(f'{k} = ?' for k in values)} WHERE id = ?",
                        list(values.values()) + [row_id],
                    )
                stats["rows_updated"] += 1
            else:
                values = dict(record, sheet_name=FLAT_SHEET_NAME, is_manually_edited=1)
                conn.execute(
                    f"INSERT INTO master_accounts ({','.join(values)}) VALUES ({','.join('?' for _ in values)})",
                    list(values.values()),
                )
                stats["rows_added"] += 1
            stats["rows_imported"] += 1

        for stats in summary.values():
            stats["reason"] = f"{stats['rows_updated']} updated, {stats['rows_added']} added; lines not in the file kept"
        conn.commit()
    finally:
        conn.close()
    return summary


def _normalize_branch(value) -> str:
    return " ".join(str(value or "").strip().lower().split())


def _normalize_service_number(value) -> Optional[str]:
    """
    Service numbers are written inconsistently across sheets — with or
    without a leading 0 or the Saudi 966 country code (966831020021512 vs
    831020021512) — so strip both for comparison purposes only.
    """
    text = _normalize_account_number(value)
    if not text:
        return None
    text = text.lstrip("0")
    if text.startswith("966") and len(text) > 9:
        text = text[3:]
    return text or None


def _accounts_compatible(master_account: Optional[str], sheet_account: str) -> bool:
    """
    True when two account numbers refer to the same account. Mobily's portal
    uses a long form (1001234009203157) where the store sheet has only the
    trailing digits (4009203157), so a suffix match counts as the same account.
    """
    if not master_account:
        return True
    return master_account == sheet_account or master_account.endswith(sheet_account) or sheet_account.endswith(master_account)


def import_store_sheet(path: str, bill_month: Optional[str] = None) -> dict:
    """
    Applies a supplementary "All stores Bills"-style spreadsheet — a single
    flat sheet of CR / Store / Service / Acc / Amount / Cost center / Dept
    columns, with no Entity/Carrier/Location of its own — onto the existing
    master_accounts rows.

    Never inserts new accounts (there's no Entity/Carrier to assign them to)
    and never deletes anything. Each sheet row is matched to one master row:
      1. by service number (the most specific key — one line = one row),
         narrowed by account number / branch name if several rows share it;
      2. otherwise by account number (exact, or suffix — see
         `_accounts_compatible`): a single-row account is unambiguous, a
         multi-row account needs a matching branch name.
    A matched row gets its CR (when numeric) and cost center from the sheet, and its
    account number filled in from the sheet's Acc when the master has none
    (e.g. the Zain sheets, which carry no account numbers). A match whose
    master account number contradicts the sheet's Acc is reported as a
    conflict and left untouched rather than guessed at.
    Updated rows are marked `is_manually_edited` so a later
    `import_master_excel` run won't wipe this enrichment.

    With `bill_month` (an archive month folder, e.g. "August_2026"), each
    row's Amount is also stored as that month's bill line — replacing any
    lines previously imported for that month — with its cost center taken
    from the sheet, else from the matched master row. Rows whose account
    isn't in the master at all (e.g. another country) are skipped.

    Returns counts plus the unmatched/conflicting sheet rows so they can be
    fixed by hand.
    """
    import openpyxl  # imported lazily: only needed for this import step

    if not os.path.isfile(path):
        raise FileNotFoundError(f"Store spreadsheet not found: {path}")

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header_map = {
        "cr": "cr", "store": "branch", "service": "service_number", "acc": "account_number",
        "cost center": "cost_center", "amount": "amount",
    }
    headers = []
    for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True)):
        key = " ".join(str(cell or "").strip().lower().split())
        headers.append(header_map.get(key))

    conn = _connect()
    result = {
        "updated": 0,
        "account_numbers_filled": 0,
        "unmatched": [],
        "conflicts": [],
        "bill_month": bill_month,
        "bill_lines_recorded": 0,
        "bill_lines_skipped": [],
    }
    try:
        _ensure_schema(conn)
        conn.row_factory = sqlite3.Row
        master_rows = conn.execute(
            "SELECT id, entity, carrier, account_number, branch, service_number, location, cost_center FROM master_accounts"
        ).fetchall()
        if bill_month:
            conn.execute("DELETE FROM bill_lines WHERE month = ?", (bill_month,))

        def record_bill_line(sheet_row: dict, record: dict, master_row, account_rows: list) -> None:
            if not bill_month:
                return
            try:
                amount = float(record.get("amount"))
            except (TypeError, ValueError):
                return
            # Which entity/carrier/account the line bills to: the matched
            # row, else the (single) master account the sheet's Acc refers to.
            owner = master_row
            if owner is None or not _accounts_compatible(owner["account_number"], sheet_row["account_number"]):
                owners = {(r["entity"], r["carrier"], r["account_number"]) for r in account_rows}
                if len(owners) != 1:
                    result["bill_lines_skipped"].append({**sheet_row, "amount": amount, "reason": "account not in master data" if not owners else "account ambiguous"})
                    return
                owner = account_rows[0]
                master_row = None
            cost_center = record.get("cost_center")
            cost_center = str(cost_center).strip() if cost_center is not None and str(cost_center).strip() else None
            if cost_center is None and master_row is not None:
                cost_center = master_row["cost_center"]
            conn.execute(
                "INSERT INTO bill_lines (month, entity, carrier, account_number, service_number, store, location, cost_center, amount, master_id, sheet_row) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                (
                    bill_month, owner["entity"], owner["carrier"], owner["account_number"] or sheet_row["account_number"],
                    sheet_row["service_number"], sheet_row["store"],
                    master_row["location"] if master_row is not None else None,
                    cost_center, amount, master_row["id"] if master_row is not None else None, sheet_row["row"],
                ),
            )
            result["bill_lines_recorded"] += 1
        by_service: dict = {}
        by_account: dict = {}
        for row in master_rows:
            service = _normalize_service_number(row["service_number"])
            if service:
                by_service.setdefault(service, []).append(row)
            if row["account_number"]:
                by_account.setdefault(row["account_number"], []).append(row)

        for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            record = {}
            for header, value in zip(headers, row):
                if header and value is not None:
                    record[header] = value

            account_number = _normalize_account_number(record.get("account_number"))
            if not account_number:
                continue
            service = _normalize_service_number(record.get("service_number"))
            branch = _normalize_branch(record.get("branch"))
            sheet_row = {
                "row": row_number,
                "account_number": account_number,
                "service_number": _normalize_account_number(record.get("service_number")),
                "store": record.get("branch"),
            }

            target = None
            same_account_rows = [
                r for acc, rows in by_account.items() if _accounts_compatible(acc, account_number) for r in rows
            ]
            candidates = by_service.get(service, []) if service else []
            if len(candidates) > 1:
                narrowed = [r for r in candidates if _accounts_compatible(r["account_number"], account_number)]
                if len(narrowed) > 1:
                    narrowed = [r for r in narrowed if _normalize_branch(r["branch"]) == branch]
                candidates = narrowed
            if len(candidates) == 1:
                target = candidates[0]
            else:
                account_rows = same_account_rows
                if len(account_rows) == 1:
                    target = account_rows[0]
                elif account_rows:
                    matches = [r for r in account_rows if _normalize_branch(r["branch"]) == branch]
                    if len(matches) == 1:
                        target = matches[0]

            if target is None:
                result["unmatched"].append(sheet_row)
                record_bill_line(sheet_row, record, None, same_account_rows)
                continue
            if not _accounts_compatible(target["account_number"], account_number):
                result["conflicts"].append({**sheet_row, "master_id": target["id"], "master_account_number": target["account_number"], "master_branch": target["branch"]})
                record_bill_line(sheet_row, record, None, same_account_rows)
                continue

            # The sheet's CR column holds the carrier name ("Mobily", "Zain")
            # instead of a CR number for non-STC lines — only take real numbers.
            cr = _normalize_account_number(record.get("cr"))
            if cr and not cr.isdigit():
                cr = None
            cost_center = record.get("cost_center")
            fill_account = None if target["account_number"] else account_number
            conn.execute(
                # A blank cell never clears a value already in the master
                # (e.g. one filled by hand), so re-importing is safe.
                "UPDATE master_accounts SET cr = COALESCE(?, cr), cost_center = COALESCE(?, cost_center), "
                "account_number = COALESCE(account_number, ?), is_manually_edited = 1 WHERE id = ?",
                (
                    cr,
                    str(cost_center).strip() if cost_center is not None else None,
                    fill_account,
                    target["id"],
                ),
            )
            result["updated"] += 1
            if fill_account:
                result["account_numbers_filled"] += 1
            filled_target = dict(target)
            filled_target["account_number"] = target["account_number"] or fill_account
            if record.get("cost_center") is None:
                filled_target["cost_center"] = target["cost_center"]
            record_bill_line(sheet_row, record, filled_target, same_account_rows)

        conn.commit()
    finally:
        conn.close()

    result["unmatched_count"] = len(result["unmatched"])
    result["conflicts_count"] = len(result["conflicts"])
    return result


UNSPECIFIED_LOCATION = "Unspecified"
MULTIPLE_LOCATIONS = "Multiple Locations"


def _resolve_location(locations: list) -> str:
    """
    One billing account number often covers many service lines (SIM cards,
    fiber links, ...) across many branches — most consistently share one
    city, but a few (15 of 317 in the source spreadsheet) genuinely span
    several. Rather than arbitrarily picking one, distinct sets collapse to
    a clearly-labeled "Multiple Locations" bucket instead of a wrong guess.
    """
    distinct = {loc for loc in locations if loc}
    if not distinct:
        return UNSPECIFIED_LOCATION
    if len(distinct) == 1:
        return next(iter(distinct))
    return MULTIPLE_LOCATIONS


def lookup_location(carrier: str, account_number: str) -> Optional[dict]:
    """
    Finds the Location/Branch/Department for one carrier+account_number.
    Location is resolved across every branch row under that account (see
    `_resolve_location`); branch/department/cr just show the first row's
    values, since those are inherently per-branch, not per-account.
    """
    if not account_number:
        return None
    conn = _connect()
    try:
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            "SELECT location, branch, department, cr FROM master_accounts WHERE carrier = ? AND account_number = ?",
            (carrier, str(account_number)),
        ).fetchall()
        if not rows:
            return None
        first = dict(rows[0])
        first["location"] = _resolve_location([r["location"] for r in rows])
        return first
    finally:
        conn.close()


def _dept_value(row) -> Optional[str]:
    """A line's department code for Oracle: Dept Code, or the department name when the code is missing."""
    for field in ("dept_code", "department"):
        value = (row[field] or "").strip() if row[field] is not None else ""
        if value:
            return value
    return None


def lookup_billing_info(entity: str, carrier: str, account_number: str) -> Optional[dict]:
    """
    Location plus every distinct cost center across all service lines under
    one entity+carrier+account — used to pick the GL charge account when
    billing that account (see oracle_sync.py). Cost centers are compared
    ignoring leading zeros, since the source sheets mix "0701" and "701".
    Returns None if the account isn't in the master at all.
    """
    if not account_number:
        return None
    conn = _connect()
    try:
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            "SELECT location, branch, cost_center, dept_code, department FROM master_accounts "
            "WHERE entity = ? AND carrier = ? AND account_number = ?",
            (entity, carrier, str(account_number)),
        ).fetchall()
    except sqlite3.OperationalError:
        return None
    finally:
        conn.close()
    if not rows:
        return None
    cost_centers = {}
    departments_by_cc: dict = {}
    for row in rows:
        cc = (row["cost_center"] or "").strip()
        cc_key = (cc.lstrip("0") or "0") if cc else None
        if cc_key:
            cost_centers.setdefault(cc_key, cc)
        dept = _dept_value(row)
        if dept and dept not in departments_by_cc.setdefault(cc_key, []):
            departments_by_cc[cc_key].append(dept)
    return {
        "location": _resolve_location([r["location"] for r in rows]),
        "branch": rows[0]["branch"],
        "line_count": len(rows),
        "lines_missing_cost_center": sum(1 for r in rows if not (r["cost_center"] or "").strip()),
        "cost_centers": sorted(cost_centers.values()),
        # Department codes per cost center (keyed like `cost_centers`, leading zeros stripped; None = no cost center).
        "departments_by_cc": departments_by_cc,
    }


def bill_line_allocation(entity: str, carrier: str, account_number: str, month: str) -> Optional[dict]:
    """
    The account's bill lines for `month`, or — when that month's sheet
    hasn't been imported — the most recent month that has lines (cost
    centers rarely change month to month, so its proportions are a sound
    basis for splitting). Returns {"month", "same_month", "lines"} or None.
    """
    conn = _connect()
    try:
        conn.row_factory = sqlite3.Row
        _ensure_schema(conn)
        rows = conn.execute(
            "SELECT b.*, m.dept_code, m.department FROM bill_lines b "
            "LEFT JOIN master_accounts m ON m.id = b.master_id "
            "WHERE b.entity = ? AND b.carrier = ? AND b.account_number = ?",
            (entity, carrier, str(account_number)),
        ).fetchall()
    finally:
        conn.close()
    if not rows:
        return None
    by_month: dict = {}
    for row in rows:
        by_month.setdefault(row["month"], []).append({**dict(row), "department_code": _dept_value(row)})
    if month in by_month:
        chosen = month
    else:
        def month_key(m):
            try:
                return datetime.strptime(m, "%B_%Y")
            except ValueError:
                return datetime.min
        chosen = max(by_month, key=month_key)
    return {"month": chosen, "same_month": chosen == month, "lines": by_month[chosen]}


def list_bill_months() -> list:
    """Months with imported bill lines and their line counts / totals."""
    conn = _connect()
    try:
        _ensure_schema(conn)
        return [
            {"month": r[0], "lines": r[1], "total": r[2]}
            for r in conn.execute("SELECT month, COUNT(*), ROUND(SUM(amount), 2) FROM bill_lines GROUP BY month ORDER BY month")
        ]
    finally:
        conn.close()


def list_master_locations() -> list:
    """Distinct cities in the master data, for mapping each to an Oracle deliver-to location."""
    conn = _connect()
    try:
        return [r[0] for r in conn.execute(
            "SELECT DISTINCT location FROM master_accounts WHERE location IS NOT NULL AND location != '' ORDER BY location"
        )]
    except sqlite3.OperationalError:
        return []
    finally:
        conn.close()


MASTER_ACCOUNT_FIELDS = [
    "entity", "carrier", "location", "cr", "branch", "department",
    "connection", "account_number", "service_number", "serial_number",
    "status", "package", "notes", "cost_center", "dept_code",
]


def list_master_accounts(
    entity: Optional[str] = None,
    carrier: Optional[str] = None,
    location: Optional[str] = None,
    search: Optional[str] = None,
    limit: int = 50,
    offset: int = 0,
) -> dict:
    """
    Lists rows from the editable master_accounts table (CRUD-backed, not
    just the read-only import), with simple filters and pagination. `search`
    matches account_number/branch/service_number/notes (case-insensitive).
    """
    conn = _connect()
    try:
        conn.row_factory = sqlite3.Row
        where = []
        params: list = []
        if entity:
            where.append("entity = ?")
            params.append(entity)
        if carrier:
            where.append("carrier = ?")
            params.append(carrier)
        if location:
            where.append("location = ?")
            params.append(location)
        if search:
            where.append("(account_number LIKE ? OR branch LIKE ? OR service_number LIKE ? OR notes LIKE ?)")
            like = f"%{search}%"
            params.extend([like, like, like, like])
        where_sql = f"WHERE {' AND '.join(where)}" if where else ""

        total = conn.execute(f"SELECT COUNT(*) FROM master_accounts {where_sql}", params).fetchone()[0]
        rows = conn.execute(
            f"SELECT * FROM master_accounts {where_sql} ORDER BY id DESC LIMIT ? OFFSET ?",
            params + [limit, offset],
        ).fetchall()
        return {"rows": [dict(r) for r in rows], "total": total}
    except sqlite3.OperationalError:
        return {"rows": [], "total": 0}
    finally:
        conn.close()


def create_master_account(fields: dict) -> int:
    """Inserts one row into master_accounts (a manual entry, not tied to any spreadsheet import). Returns its id."""
    conn = _connect()
    try:
        _ensure_schema(conn)
        values = {k: fields.get(k) for k in MASTER_ACCOUNT_FIELDS}
        values["sheet_name"] = "Manual Entry"
        values["is_manually_edited"] = 1
        columns = list(values.keys())
        placeholders = ",".join("?" for _ in columns)
        cur = conn.execute(
            f"INSERT INTO master_accounts ({','.join(columns)}) VALUES ({placeholders})",
            [values[c] for c in columns],
        )
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def update_master_account(row_id: int, fields: dict) -> bool:
    """
    Updates the given fields on one master_accounts row, and marks it as
    manually edited so a future spreadsheet re-import won't overwrite it
    (see `import_master_excel`). Returns False if the row doesn't exist.
    """
    updates = {k: v for k, v in fields.items() if k in MASTER_ACCOUNT_FIELDS}
    if not updates:
        return False
    updates["is_manually_edited"] = 1
    conn = _connect()
    try:
        _ensure_schema(conn)
        set_sql = ", ".join(f"{k} = ?" for k in updates)
        cur = conn.execute(
            f"UPDATE master_accounts SET {set_sql} WHERE id = ?",
            list(updates.values()) + [row_id],
        )
        conn.commit()
        return cur.rowcount > 0
    finally:
        conn.close()


def delete_master_account(row_id: int) -> bool:
    conn = _connect()
    try:
        cur = conn.execute("DELETE FROM master_accounts WHERE id = ?", (row_id,))
        conn.commit()
        return cur.rowcount > 0
    finally:
        conn.close()


def build_location_index() -> dict:
    """
    Loads the whole (carrier, account_number) -> location info mapping into
    memory in one query, for bulk-joining against a list of accounts/invoices
    without one SQLite query per row. Location is resolved across every
    branch row under that account (see `_resolve_location`); branch/
    department just show the first row's values, since those are
    inherently per-branch, not per-account.
    """
    conn = _connect()
    index = {}
    grouped = {}
    try:
        if not os.path.isfile(DB_PATH):
            return index
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            "SELECT carrier, account_number, location, branch, department FROM master_accounts"
        ).fetchall()
        for row in rows:
            key = (row["carrier"], row["account_number"])
            grouped.setdefault(key, []).append(row)

        for key, group_rows in grouped.items():
            first = group_rows[0]
            index[key] = {
                "location": _resolve_location([r["location"] for r in group_rows]),
                "branch": first["branch"],
                "department": first["department"],
            }
        return index
    except sqlite3.OperationalError:
        return index
    finally:
        conn.close()
